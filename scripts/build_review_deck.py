"""Build an editable review deck from verified CSVs; optional python-pptx dependency."""
import csv
from pathlib import Path
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE,XL_LEGEND_POSITION
from pptx.dml.color import RGBColor
from pptx.util import Inches,Pt
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.xmlchemy import OxmlElement

ROOT=Path(__file__).resolve().parents[1]
R=Presentation();R.slide_width=Inches(13.333);R.slide_height=Inches(7.5)
R.core_properties.author="Perplexity Computer"
R.core_properties.title="UPI Fraud Stress Lab: Faculty Review"
BG="171C1E";FG="ECEBE6";MUTED="B2BEBD";TEAL="80B8A5";WARN="E2B25B"
def rgb(v):return RGBColor.from_string(v)
def text(s,value,x,y,w,h,size=20,color=FG,bold=False):
    box=s.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h))
    tf=box.text_frame;tf.word_wrap=True
    tf.margin_left=tf.margin_right=0;tf.margin_top=tf.margin_bottom=0
    for i,line in enumerate(value.split("\n")):
        p=tf.paragraphs[0] if i==0 else tf.add_paragraph();p.text=line
        p.font.name="Liberation Sans";p.font.size=Pt(size);p.font.bold=bold;p.font.color.rgb=rgb(color)
        p.space_after=Pt(12)
    return box
def slide(title,source="Project records: results/ and docs/VERIFICATION_REPORT.md"):
    s=R.slides.add_slide(R.slide_layouts[6]);s.background.fill.solid();s.background.fill.fore_color.rgb=rgb(BG)
    text(s,title,.6,.48,12,1.05,34,bold=True)
    text(s,"SYNTHETIC RESEARCH  /  GATE C HOLD",.6,6.93,9,.3,11,TEAL)
    text(s,str(len(R.slides)),12.1,6.93,.5,.3,11,MUTED)
    text(s,source,.6,6.45,12,.32,11,MUTED)
    return s
def block(s,title,body,x,y,w=5.6):
    text(s,title,x,y,w,.45,23,TEAL,True);text(s,body,x,y+.65,w,2.8,20)
def chart(s,categories,series,x,y,w,h,axis_title):
    data=CategoryChartData();data.categories=categories
    for name,values in series:data.add_series(name,values)
    c=s.shapes.add_chart(XL_CHART_TYPE.BAR_CLUSTERED,Inches(x),Inches(y),Inches(w),Inches(h),data).chart
    c.has_legend=len(series)>1
    if c.has_legend:c.legend.position=XL_LEGEND_POSITION.BOTTOM;c.legend.font.size=Pt(13);c.legend.font.color.rgb=rgb(FG)
    c.chart_style=10
    c.has_title=False
    sp=OxmlElement("c:spPr");fill=OxmlElement("a:solidFill");color=OxmlElement("a:srgbClr")
    color.set("val",BG);fill.append(color);sp.append(fill)
    c._chartSpace.insert(list(c._chartSpace).index(c._chartSpace.chart)+1,sp)
    for axis in (c.category_axis,c.value_axis):
        axis.tick_labels.font.name="Liberation Sans";axis.tick_labels.font.size=Pt(13);axis.tick_labels.font.color.rgb=rgb(FG)
    c.value_axis.minimum_scale=0;c.value_axis.maximum_scale=100;c.value_axis.major_unit=25
    c.value_axis.has_title=True;c.value_axis.axis_title.text_frame.text=axis_title
    for p in c.value_axis.axis_title.text_frame.paragraphs:p.font.size=Pt(13);p.font.color.rgb=rgb(FG)
    for n,ser in enumerate(c.series):
        ser.format.fill.solid();ser.format.fill.fore_color.rgb=rgb(TEAL if n==0 else WARN)
    return c

rows=list(csv.DictReader((ROOT/"results/results_summary.csv").open()))
s=slide("UPI Fraud Stress Lab")
text(s,"Campaign-level evaluation before intervention",.6,1.8,11,1,30,TEAL)
text(s,"An auditable engineering case study, not a deployed UPI detector.\nWorking software, matched synthetic controls and preserved adverse results.",.6,3.1,11,1.8,25)
text(s,"Academic review draft • author, venue and external checker approval pending",.6,5.5,12,.5,16,MUTED)

s=slide("Why evaluate a campaign?")
block(s,"Transaction score","AUC ranks examples in a specified dataset.\nIt does not say when a sequence stops.",.6,1.9)
block(s,"Campaign endpoint","Does a simulated actor reach its target before BLOCK?\nFLAG can still transfer; BLOCK terminates.",7,1.9)
text(s,"Gross simulated debit is not victim loss or proven savings.",.6,5.4,12,.6,24,WARN)

s=slide("The implemented evidence pipeline")
labels=["Local dashboard","Safe Python runner","Existing engines","Immutable records"]
for i,label in enumerate(labels):
    x=.6+i*3.15
    shape=s.shapes.add_shape(MSO_SHAPE.RECTANGLE,Inches(x),Inches(2.1),Inches(2.7),Inches(1.1))
    shape.fill.solid();shape.fill.fore_color.rgb=rgb("253235");shape.line.color.rgb=rgb("465351")
    text(s,label,x+.15,2.44,2.4,.4,19)
    if i<3:text(s,"→",x+2.8,2.4,.35,.4,22,TEAL)
text(s,"Independent checkers and raw-record reconstruction feed tables, figures and the paper.\nNo arbitrary shell, public binding or Git/Gate C endpoint.",.6,4.2,12,1.35,23)

s=slide("A narrow, registered exploratory pilot")
block(s,"Data and detectors","9 independent blocks; 18 fresh detector fits.\n7,200 evaluated attack campaigns.\nHVB allowed; declared GE excluded from development.\nOrdinary benign and GE-mimic test strata.",.6,1.8)
block(s,"Five online-matched controls","Static; rule-adaptive; trained RL;\nexplicit-label-blind RL; zero Q.\nCommon actions, clipping, pacing and budgets.\nOffline training costs remain unequal.",7,1.8)
text(s,"No powered confirmatory lockbox and no real-world UPI validation.",.6,5.7,12,.5,22,WARN)

s=slide("Completion depends on detector and control")
arms=["Static","Rule","RL","Blind RL","Zero Q"];keys=["static","rule","rl","blind_rl","zero_q"]
vals=lambda d:[next(float(r["completion_pct"]) for r in rows if r["detector"]==d and r["family"]=="gradual_escalation" and r["strategy"]==a) for a in keys]
chart(s,arms,[("XGBoost",vals("xgboost")),("Random Forest",vals("random_forest"))],.55,1.75,8.2,4.25,"Completion (%)")
text(s,"Held-out GE\n360 campaigns per arm/detector\n9 independent blocks",9.1,1.9,3.5,1.5,18,TEAL)
sd=lambda d:[next(float(r["block_sd_pp"]) for r in rows if r["detector"]==d and r["family"]=="gradual_escalation" and r["strategy"]==a) for a in keys]
text(s,"Block SD (percentage points)\nArm: XGB / RF\n"+"\n".join(f"{a}: {x:.2f} / {y:.2f}" for a,x,y in zip(arms,sd("xgboost"),sd("random_forest"))),9.1,3.55,3.6,2.35,15)

s=slide("Benign mechanism shift is a major failure")
chart(s,["XGBoost","Random Forest"],[("Legitimate GE mimics",[1738/36,461/36])],.55,1.8,7.9,3.8,"Benign false-block (%)")
text(s,"48.278%\n1,738 / 3,600 XGBoost episodes\n\n12.806%\n461 / 3,600 RF episodes",8.9,1.85,3.8,3.5,23,WARN)
text(s,"Ordinary benign: 0 observed BLOCK / 3,600 per detector.\nNo population guarantee; nine blocks. Mechanism exclusion ≠ reliable discrimination.",.6,5.45,12,1,19)

s=slide("What we can defensibly contribute",source="Source: verified literature links below; project implementation and retained results")
block(s,"Prior research already does","Sequence fraud detection, synthetic finance,\nbanking evasion and learned attackers.\nThis project is not the first RL fraud attack.",.6,1.8)
block(s,"Implemented contribution","Auditable campaign/ledger infrastructure,\nfive matched controls, narrow lineage checks,\nbenign-shift controls and reconstruction.\nNo universal superiority claim.",7,1.8)
for x,label,url in [( .6,"Fraud-RLA","https://doi.org/10.1109/TDSC.2026.3673847"),
                    (4.5,"Banking evasion","https://www.usenix.org/conference/raid2020/presentation/carminati"),
                    (8.6,"PaySim","https://www.msc-les.org/proceedings/emss/2016/EMSS2016_249.pdf")]:
    b=text(s,label,x,5.45,3.5,.4,17,TEAL)
    b.click_action.hyperlink.address=url

s=slide("A real local research workstation")
from PIL import Image
path=ROOT/"docs/images/dashboard.png";crop=ROOT/"docs/images/dashboard-review-crop.png"
with Image.open(path) as im:
    im.crop((260,310,min(im.width,1490),710)).save(crop)
    aspect=(min(im.width,1490)-260)/400
w=12.1;h=w/aspect
s.shapes.add_picture(str(crop),Inches(.6),Inches(1.65),width=Inches(w),height=Inches(h))
text(s,"Actual overview detail. Live demo: run → logs → campaign → comparisons → export.",.6,5.95,12,.35,16,TEAL)

s=slide("Verified engineering; bounded scientific claims")
block(s,"Executed verification","478 distinct tests passed in fresh Python 3.11.\nShort legacy/pilot main, repeat, reload and check.\n1,012,673 retained-pilot verifier checks;\n30,985 TD updates replayed.",.6,1.75)
block(s,"Still requires human action","Approve authors, venue and AI-use policy.\nObtain genuine approved AI/similarity reports.\nSelect a project license.\nDo not infer real-world effectiveness.",7,1.75)
text(s,"Ready for exploratory academic review, not unqualified submission.\nScientific validity extends only as far as the evidence. Gate C remains HOLD.",.6,5.25,12,1,21,WARN)
out=ROOT/"research/review_materials/review_deck.pptx";out.parent.mkdir(parents=True,exist_ok=True);R.save(out)
print(out)
